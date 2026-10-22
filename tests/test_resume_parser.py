"""
Unit Tests for Resume Parser Module
Tests PDF text extraction, DOCX text extraction, empty files,
corrupted files, and text normalization.
"""

import io
import pytest
import docx  # python-docx

from modules.resume_parser import (
    extract_resume_text,
    extract_text_from_pdf,
    extract_text_from_docx,
    normalize_extracted_text,
    validate_file_metadata
)


class TestResumeParser:
    def test_text_normalization(self):
        """
        Verifies that extraneous spaces and control characters are cleaned
        while preserving paragraphs.
        """
        raw = "John   Doe  \r\n\r\nSoftware   Engineer\n\n\n\nSkills: Python,    Docker\x00"
        cleaned = normalize_extracted_text(raw)
        assert "\r" not in cleaned
        assert "\x00" not in cleaned
        assert "John Doe" in cleaned
        assert "Software Engineer" in cleaned
        assert "Skills: Python, Docker" in cleaned

    def test_empty_file_handling(self):
        """
        Verifies that an empty (0 byte) file is rejected with clear error.
        """
        is_valid, err = validate_file_metadata("resume.pdf", 0)
        assert not is_valid
        assert "empty" in err.lower()

        result = extract_resume_text(("empty.pdf", b""))
        assert result["success"] is False
        assert "empty" in result["error"].lower()

    def test_unsupported_file_extension(self):
        """
        Verifies that unsupported file extensions are rejected.
        """
        is_valid, err = validate_file_metadata("malicious.exe", 1024)
        assert not is_valid
        assert "unsupported" in err.lower()

        result = extract_resume_text(("script.sh", b"echo 'test'"))
        assert result["success"] is False

    def test_pdf_extraction(self):
        """
        Creates an in-memory PDF and verifies clean extraction.
        """
        from reportlab.pdfgen import canvas
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer)
        c.drawString(50, 750, "Alex Morgan")
        c.drawString(50, 730, "Email: alex.morgan@example.com | Phone: (555) 234-5678")
        c.drawString(50, 710, "Summary: Senior Backend Engineer with 5+ years of experience in Python and Cloud systems.")
        c.drawString(50, 690, "Skills: Python, Django, PostgreSQL, Docker, Kubernetes, AWS")
        c.drawString(50, 670, "Experience: Senior Developer at TechCorp (2021 - Present). Architected REST APIs.")
        c.showPage()
        c.save()
        pdf_bytes = buffer.getvalue()

        result = extract_text_from_pdf(pdf_bytes, "alex_morgan.pdf")
        assert result["success"] is True
        assert "Alex Morgan" in result["text"]
        assert "alex.morgan@example.com" in result["text"]
        assert "Python" in result["text"]
        assert result["word_count"] > 10
        assert result["page_count"] == 1
        assert result["is_scanned"] is False

    def test_docx_extraction_with_tables(self):
        """
        Creates an in-memory DOCX with paragraphs and tables using python-docx
        and verifies extraction.
        """
        doc = docx.Document()
        doc.add_heading("Sarah Connor", level=1)
        doc.add_paragraph("Full Stack Developer specializing in React and Node.js.")
        
        # Add table
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Frontend"
        table.cell(0, 1).text = "React, TypeScript"
        table.cell(1, 0).text = "Backend"
        table.cell(1, 1).text = "Node.js, Express"

        doc_stream = io.BytesIO()
        doc.save(doc_stream)
        docx_bytes = doc_stream.getvalue()

        result = extract_text_from_docx(docx_bytes, "sarah_connor.docx")
        assert result["success"] is True
        assert "Sarah Connor" in result["text"]
        assert "React" in result["text"]
        assert "TypeScript" in result["text"]
        assert "Node.js" in result["text"]
        assert result["word_count"] > 10

    def test_corrupted_file_handling(self):
        """
        Verifies that corrupted bytes in a .pdf or .docx file return a friendly error.
        """
        corrupted_bytes = b"%PDF-1.4\nJUNK CORRUPTED BYTES RANDOM DATA"
        result_pdf = extract_text_from_pdf(corrupted_bytes, "corrupt.pdf")
        assert result_pdf["success"] is False
        assert result_pdf["error"] is not None

        result_docx = extract_text_from_docx(b"NOT A ZIP OR DOCX ARCHIVE", "corrupt.docx")
        assert result_docx["success"] is False
        assert result_docx["error"] is not None
